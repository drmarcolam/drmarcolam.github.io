import openpyxl, json, re, sys, docx
SRC = '/tmp/v2out/Assignment_1_V2_Instructor_Completed.xlsx'
wb = openpyxl.load_workbook(SRC, data_only=True)

def entries(sheet):
    ws = wb[sheet]; out = {}; order = []
    for r in ws.iter_rows(min_row=1, values_only=True):
        e, acct, dr, cr, ex = r[0], r[1], r[2], r[3], r[4]
        if not e or e == 'Entry' or not acct: continue
        if not re.match(r'^((GF|SR)-\d+|[AC]\d+)$', str(e)): continue
        if e not in out: out[e] = []; order.append(e)
        out[e].append({'a': acct, 'dr': float(dr or 0), 'cr': float(cr or 0), 'x': ex or ''})
    return [{'id': k, 'lines': out[k]} for k in order]

def coa(sheet):
    ws = wb[sheet]; res = []
    for r in ws.iter_rows(min_row=4, values_only=True):
        if r[0] and r[0] != 'Account':
            res.append({'a': r[0], 'bd': float(r[1] or 0), 'bc': float(r[2] or 0)})
    return res

# transaction text
ws = wb['Transactions']; tx = {}; adj = {}; mode = 'tx'
for r in ws.iter_rows(values_only=True):
    if r[0] and str(r[0]).startswith('Year-End'): mode = 'adj'; continue
    if r[0] and r[1] and re.match(r'^(GF|SR)-\d+$', str(r[0])): tx[r[0]] = r[1]
    if r[0] and r[1] and re.match(r'^A\d+$', str(r[0])): adj[r[0]] = r[1]

funds = []
for key, name, prefix, jsheet, asheet, csheet, ccoa in [
    ('GF', 'General Fund', 'GF', 'GF Journal', 'GF Adjustments', 'GF Closing Entries', 'GF Chart of Accounts'),
    ('SRF', 'Public Safety Grant Special Revenue Fund', 'SR', 'SRF Journal', 'SRF Adjustments', 'SRF Closing Entries', 'SRF Chart of Accounts')]:
    funds.append({'key': key, 'name': name, 'coa': coa(ccoa), 'journal': entries(jsheet), 'adjust': entries(asheet), 'closing': entries(csheet)})

# transactions without entries (info only)
for f in funds:
    ids = [e['id'] for e in f['journal']]
    pref = 'GF-' if f['key'] == 'GF' else 'SR-'
    f['tx'] = [{'id': k, 'text': v, 'hasEntry': k in ids} for k, v in tx.items() if k.startswith(pref)]
    have = {e['id'] for e in f['adjust']}
    ids = ['A%d' % i for i in range(1, 8)] if f['key'] == 'GF' else ['A8', 'A9']
    f['adjList'] = [{'id': i, 'text': adj[i], 'hasEntry': i in have} for i in ids]

# checkpoints from solution docx
d = docx.Document('/mnt/user-data/uploads/Accounting Case/Assignment_1_V2_Instructor_Solution.docx')
kt = [t for t in d.tables if [c.text for c in t.rows[0].cells] == ['Item', 'Amount']]
def money(s): return float(s.replace('$', '').replace(',', '').replace('(', '-').replace(')', ''))
cps = {'GF': [(r.cells[0].text, money(r.cells[1].text)) for r in kt[0].rows[1:]],
       'SRF': [(r.cells[0].text, money(r.cells[1].text)) for r in kt[1].rows[1:]]}
for f in funds: f['checkpoints'] = [{'label': l, 'v': v} for l, v in cps[f['key']]]

# verify: replay key against recalculated adjusted TB
def replay(f, sheetTB):
    bal = {c['a']: c['bd'] - c['bc'] for c in f['coa']}
    for grp in ('journal', 'adjust'):
        for e in f[grp]:
            for l in e['lines']: bal[l['a']] = bal.get(l['a'], 0) + l['dr'] - l['cr']
    ws = wb[sheetTB]; bad = []
    for r in ws.iter_rows(min_row=5, values_only=True):
        if r[0] and r[0] in bal and isinstance(r[-1], (int, float)):
            if abs(bal[r[0]] - r[-1]) > 0.5: bad.append((r[0], bal[r[0]], r[-1]))
    return bad, bal
for f, tb in zip(funds, ['GF Adjusted TB', 'SRF Adjusted TB']):
    bad, bal = replay(f, tb)
    print(f['key'], 'journal', len(f['journal']), 'adj', [e['id'] for e in f['adjust']], 'closing', [e['id'] for e in f['closing']], 'tx', len(f['tx']), 'mismatch', bad)
    print('   tx w/o entry:', [t['id'] for t in f['tx'] if not t['hasEntry']])
    # closing leaves unassigned FB
    for e in f['closing']:
        for l in e['lines']: bal[l['a']] = bal.get(l['a'], 0) + l['dr'] - l['cr']
    print('   post-closing nonzero:', {k: v for k, v in bal.items() if abs(v) > 0.5 and not k.startswith(('Encumb',))})
    print('   checkpoints', f['checkpoints'])
LAB = {'GF': {'C1': 'Close the budgetary accounts (estimated revenues, appropriations, estimated other financing uses and the budgetary fund balance).', 'C2': 'Close the revenue accounts to Fund Balance - Unassigned.', 'C3': 'Close the expenditure accounts to Fund Balance - Unassigned.', 'C4': 'Close the transfers out to Fund Balance - Unassigned.'},
       'SRF': {'C1': 'Close revenues, the transfer in and expenditures, and record the ending restricted fund balance.'}}
SPEC = {'GF': {'Ending fund balance': {'t': 'fb'}, 'Cash': {'t': 'acct', 'a': 'Cash', 'side': 'dr'}, 'Due to Internal Service Fund': {'t': 'acct', 'a': 'Due to Internal Service Fund', 'side': 'cr'},
               'Property tax revenue': {'t': 'acct', 'a': 'Revenues - Property Taxes', 'side': 'cr'}, 'Total revenues': {'t': 'sum', 'p': 'Revenues - ', 'side': 'cr'}, 'Total expenditures': {'t': 'sum', 'p': 'Expenditures - ', 'side': 'dr'}},
        'SRF': {'Ending fund balance': {'t': 'fb'}, 'Grant revenue': {'t': 'acct', 'a': 'Revenues - Intergovernmental Grant', 'side': 'cr'}, 'Program expenditures': {'t': 'acct', 'a': 'Expenditures - Public Safety Program', 'side': 'dr'},
                'Grant receivable': {'t': 'acct', 'a': 'Grant Receivable', 'side': 'dr'}, 'Cash': {'t': 'acct', 'a': 'Cash', 'side': 'dr'}}}
for f in funds:
    for e in f['closing']: e['label'] = LAB[f['key']][e['id']]
    for c in f['checkpoints']: c['spec'] = SPEC[f['key']][c['label']]
import json as _j
obj = {'id': 'a1', 'title': 'Assignment 1: Operating Governmental Funds', 'funds': funds}
open('/home/claude/build/v2/a1.json', 'w').write(_j.dumps(obj, indent=1))
open('/home/claude/build/v2/a1.js', 'w').write('window.CASE = ' + _j.dumps(obj, separators=(',', ':')) + ';\n')
print('written')
json.dump({'id': 'a1', 'title': 'Assignment 1: Operating Governmental Funds', 'funds': funds}, open('/home/claude/build/v2/a1.json', 'w'), indent=1)
print('adj all', sorted(adj))
