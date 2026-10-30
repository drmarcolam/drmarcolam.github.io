import openpyxl, json, re, docx, sys
UP = '/mnt/user-data/uploads/Accounting Case/'
OUT = '/home/claude/build/v2/'
def money(s): return float(s.replace('$', '').replace(',', '').replace('(', '-').replace(')', ''))
ID = re.compile(r'^([A-Z]+-\d+|[AC]\d+)$')

CFG = {
 'a2': dict(n=2, title='Assignment 2: Capital Projects and Debt Service Funds', funds=[
   dict(key='CPF', name='Capital Projects Fund (Community Recreation Center)', sh='CPF', tx='CP-', adj=['A1'], drop=[],
        close='Close revenues, other financing sources and expenditures, and record the ending restricted fund balance.',
        cps={'Ending restricted fund balance': dict(t='fb', p='Fund Balance'), 'Cash': dict(t='acct', a='Cash', side='dr'), 'Contracts payable': dict(t='acct', a='Contracts Payable', side='cr'),
             'Retainage payable': dict(t='acct', a='Retainage Payable', side='cr'), 'Capital outlay': dict(t='acct', a='Expenditures - Capital Outlay', side='dr'), 'Outstanding encumbrance': dict(t='acct', a='Encumbrances Outstanding', side='cr')}),
   dict(key='DSF', name='Debt Service Fund', sh='DSF', tx='DS-', adj=['A2'], drop=[],
        close='Close revenues, the transfer in and expenditures, and record the change in restricted fund balance.',
        cps={'Ending restricted fund balance': dict(t='fb', p='Fund Balance'), 'Cash': dict(t='acct', a='Cash', side='dr'), 'Principal expenditure': dict(t='acct', a='Expenditures - Bond Principal', side='dr'), 'Interest expenditure': dict(t='acct', a='Expenditures - Bond Interest', side='dr')})]),
 'a3': dict(n=3, title='Assignment 3: Proprietary Funds', funds=[
   dict(key='ENT', name='Water Utility Enterprise Fund', sh='Enterprise', tx='E-', adj=['A1'], drop=['E-14'],
        close='Close the revenue and expense accounts to Net Position.',
        cps={'Ending net position': dict(t='fb', p='Net Position'), 'Change in net position': dict(t='chg', p='Net Position'), 'Cash': dict(t='acct', a='Cash', side='dr'),
             'Accounts receivable, net': dict(t='accts', a=['Accounts Receivable', 'Allowance for Doubtful Accounts'], side='dr'), 'Capital assets, net': dict(t='accts', a=['Capital Assets', 'Accumulated Depreciation'], side='dr'), 'Revenue bonds payable': dict(t='acct', a='Revenue Bonds Payable', side='cr')}),
   dict(key='ISF', name='Central Motor Pool Internal Service Fund', sh='ISF', tx='I-', adj=['A2', 'A3'], drop=['I-8'],
        close='Close the revenue and expense accounts to Net Position.',
        cps={'Ending net position': dict(t='fb', p='Net Position'), 'Current-year loss': dict(t='chg', p='Net Position'), 'Due from General Fund': dict(t='acct', a='Due from General Fund', side='dr'), 'Cash': dict(t='acct', a='Cash', side='dr'),
             'Capital assets, net': dict(t='accts', a=['Capital Assets', 'Accumulated Depreciation'], side='dr')})]),
 'a4': dict(n=4, title='Assignment 4: Custodial Fiduciary Fund', funds=[
   dict(key='CUS', name='Regional Tax Collection Custodial Fund', sh='Custodial', tx='F-', adj=['A1', 'A2'], drop=[],
        close='Close additions and deductions. Fiduciary net position ends at zero.',
        cps={'Ending cash': dict(t='acct', a='Cash', side='dr'), 'Due to county': dict(t='acct', a='Due to County', side='cr'), 'Due to school district': dict(t='acct', a='Due to School District', side='cr'),
             'Ending fiduciary net position': dict(t='fb', p='Fiduciary Net Position'), 'Additions': dict(t='sum', p='Additions - ', side='cr'), 'Deductions': dict(t='sum', p='Deductions - ', side='dr')})]),
}

def build(cid, cfg):
    n = cfg['n']
    wb = openpyxl.load_workbook('/tmp/v2out/Assignment_%d_V2_Instructor_Completed.xlsx' % n, data_only=True)
    d = docx.Document(UP + 'Assignment_%d_V2_Instructor_Solution.docx' % n)
    kts = [t for t in d.tables if [c.text for c in t.rows[0].cells] == ['Item', 'Amount']]
    tx = {}; adj = {}; mode = 'tx'
    for r in wb['Transactions'].iter_rows(values_only=True):
        if r[0] and r[1] and ID.match(str(r[0])):
            (adj if re.match(r'^A\d+$', str(r[0])) else tx)[r[0]] = r[1]
    def entries(sheet):
        out = {}; order = []
        for r in wb[sheet].iter_rows(values_only=True):
            e, a = r[0], r[1]
            if not e or e == 'Entry' or not a or not ID.match(str(e)): continue
            if e not in out: out[e] = []; order.append(e)
            out[e].append({'a': a, 'dr': float(r[2] or 0), 'cr': float(r[3] or 0), 'x': r[4] or ''})
        return [{'id': k, 'lines': out[k]} for k in order]
    funds = []
    for fi, fc in enumerate(cfg['funds']):
        sh = fc['sh']
        coa = [{'a': r[0], 'bd': float(r[1] or 0), 'bc': float(r[2] or 0)} for r in wb[sh + ' Chart of Accounts'].iter_rows(min_row=4, values_only=True) if r[0] and r[0] != 'Account']
        f = {'key': fc['key'], 'name': fc['name'], 'coa': coa, 'journal': entries(sh + ' Journal'), 'adjust': entries(sh + ' Adjustments'), 'closing': entries(sh + ' Closing Entries')}
        have = {e['id'] for e in f['journal']}
        f['tx'] = [{'id': k, 'text': v, 'hasEntry': k in have} for k, v in tx.items() if k.startswith(fc['tx']) and k not in fc['drop']]
        ah = {e['id'] for e in f['adjust']}
        f['adjList'] = [{'id': i, 'text': adj[i], 'hasEntry': i in ah} for i in fc['adj']]
        for e in f['closing']: e['label'] = fc['close']
        kt = kts[fi]
        f['checkpoints'] = [{'label': r.cells[0].text, 'v': money(r.cells[1].text), 'spec': fc['cps'][r.cells[0].text]} for r in kt.rows[1:]]
        funds.append(f)
    return {'id': cid, 'title': cfg['title'], 'funds': funds}

def spec_val(sp, tb, coa):
    def sgn(a): return tb.get(a, (0, 0))[0]
    if sp['t'] == 'acct': x = sgn(sp['a']); return x if sp['side'] == 'dr' else -x
    if sp['t'] == 'accts': x = sum(sgn(a) for a in sp['a']); return x if sp['side'] == 'dr' else -x
    if sp['t'] == 'sum':
        x = sum(v[0] for k, v in tb.items() if k.startswith(sp['p'])); return x if sp['side'] == 'dr' else -x
    if sp['t'] == 'fb': return -sum(v[1] for k, v in tb.items() if k.startswith(sp['p']))
    if sp['t'] == 'chg': return -sum(v[1] for k, v in tb.items() if k.startswith(sp['p'])) + sum(c['bd'] - c['bc'] for c in coa if c['a'].startswith(sp['p']))

ok = True
for cid, cfg in CFG.items():
    case = build(cid, cfg)
    wb = openpyxl.load_workbook('/tmp/v2out/Assignment_%d_V2_Instructor_Completed.xlsx' % cfg['n'], data_only=True)
    for f, fc in zip(case['funds'], cfg['funds']):
        bal = {c['a']: c['bd'] - c['bc'] for c in f['coa']}
        for g in ('journal', 'adjust'):
            for e in f[g]:
                for l in e['lines']: bal[l['a']] = bal.get(l['a'], 0) + l['dr'] - l['cr']
        bad = []
        for r in wb[fc['sh'] + ' Adjusted TB'].iter_rows(min_row=5, values_only=True):
            if r[0] in bal and isinstance(r[-1], (int, float)) and abs(bal[r[0]] - r[-1]) > 0.5: bad.append((r[0], bal[r[0]], r[-1]))
        adjb = dict(bal); post = dict(bal)
        for e in f['closing']:
            for l in e['lines']: post[l['a']] = post.get(l['a'], 0) + l['dr'] - l['cr']
        tb = {a: (adjb[a], post[a]) for a in adjb}
        cpbad = []
        for c in f['checkpoints']:
            v = spec_val(c['spec'], tb, f['coa'])
            if abs(v - c['v']) > 0.5: cpbad.append((c['label'], v, c['v']))
        resid = {k: v for k, v in post.items() if abs(v) > 0.5 and (k.startswith(('Revenues', 'Expenditures', 'Operating', 'Nonoperating', 'Additions', 'Deductions', 'Other Financing')))}
        print(cid, f['key'], 'journal', len(f['journal']), 'adj', [e['id'] for e in f['adjust']], 'info adj', [a['id'] for a in f['adjList'] if not a['hasEntry']], 'closing', len(f['closing']), 'tx', len(f['tx']),
              '| TB mismatch', bad, '| cp mismatch', cpbad, '| unclosed temp accts', resid)
        if bad or cpbad or resid: ok = False
    open(OUT + cid + '.js', 'w').write('window.CASE = ' + json.dumps(case, separators=(',', ':')) + ';\n')
print('ALL OK' if ok else 'PROBLEMS')
