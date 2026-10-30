import openpyxl, json, re
wb = openpyxl.load_workbook('/tmp/v2out/Assignment_5_V2_Instructor_Completed.xlsx', data_only=True)
def rows(sh, minr=4):
    out = []
    for r in wb[sh].iter_rows(min_row=minr, values_only=True):
        if any(c is not None for c in r): out.append(r)
    return out
def find(sh, label):
    for r in wb[sh].iter_rows(values_only=True):
        if r[0] == label: return r
    raise KeyError(label)
facts = []
for r in wb['Transactions'].iter_rows(values_only=True):
    if r[0] and r[1] and re.match(r'^GW-\d+$', str(r[0])): facts.append({'id': r[0], 'text': r[1]})
S = []
# S1 fund-level results
fc = [('General Fund ending fund balance', 'fb_gf'), ('Special Revenue Fund ending fund balance', 'fb_srf'), ('Capital Projects Fund ending fund balance', 'fb_cpf'), ('Debt Service Fund ending fund balance', 'fb_dsf'),
      ('Total governmental fund balances', 'fb_tot'), ('Enterprise Fund ending net position', 'np_ent'), ('Internal Service Fund ending net position', 'np_isf'), ('Custodial Fund fiduciary net position', 'np_cus')]
S.append({'id': 's1', 'title': 'Part 1: Fund-level results carried forward', 'intro': 'Bring forward the ending results from Assignments 1 to 4. Enter 0 where the amount is zero.', 'cols': ['Item', 'Amount'],
          'rows': [{'label': l, 'cells': [{'k': k, 'a': float(find('Fund Checkpoints', l)[1])}]} for l, k in fc]})
# S2 conversion data
cap = find('Net Position Reconciliation', 'Add governmental capital assets, net')[1]
bonds = -find('Net Position Reconciliation', 'Less governmental bonds payable')[1]
taxadj = find('Change Reconciliation', 'Property-tax timing difference')[1]
S.append({'id': 's2', 'title': 'Part 2: Government-wide conversion data', 'intro': 'Use the case facts above to compute these amounts. Enter decreases as negative numbers (for example -103000 or (103000)).', 'cols': ['Item', 'Amount'], 'rows': [
  {'label': 'Governmental capital assets, net, at June 30 (beginning + capital outlay and equipment - depreciation)', 'cells': [{'k': 'cap', 'a': float(cap)}]},
  {'label': 'Governmental bonds payable at June 30 (beginning + new bonds - principal paid)', 'cells': [{'k': 'bonds', 'a': float(bonds)}]},
  {'label': 'Property-tax availability adjustment to the change in net position', 'cells': [{'k': 'taxadj', 'a': float(taxadj)}]},
  {'label': 'Net change in fund balance: General Fund', 'cells': [{'k': 'cf_gf', 'a': 1077000.0}]},
  {'label': 'Net change in fund balance: Special Revenue Fund', 'cells': [{'k': 'cf_srf', 'a': 60000.0}]},
  {'label': 'Net change in fund balance: Capital Projects Fund', 'cells': [{'k': 'cf_cpf', 'a': 930000.0}]},
  {'label': 'Net change in fund balance: Debt Service Fund', 'cells': [{'k': 'cf_dsf', 'a': -80000.0}]},
  {'label': 'Change in total governmental fund balances', 'cells': [{'k': 'cf_tot', 'a': float(find('Change Reconciliation', 'Change in total governmental fund balances')[1])}]}]})
# S3 NP recon
npl = [('Total governmental fund balances', 'np1'), ('Add governmental capital assets, net', 'np2'), ('Less governmental bonds payable', 'np3'), ('Less compensated absences', 'np4'),
       ('Add unavailable property taxes recognized government-wide', 'np5'), ('Add Internal Service Fund net position', 'np6'), ('Governmental activities net position', 'np_tot')]
S.append({'id': 's3', 'title': 'Part 3: Reconcile fund balances to net position', 'intro': 'Governmental funds to governmental activities, ending balance sheet. Enter deductions as negative numbers.', 'cols': ['Reconciliation item', 'Amount'],
          'rows': [{'label': l, 'cells': [{'k': k, 'a': float(find('Net Position Reconciliation', l)[1])}]} for l, k in npl]})
# S4 change recon
chl = [('Change in total governmental fund balances', 'ch1'), ('Capital outlays capitalized', 'ch2'), ('Depreciation expense', 'ch3'), ('Bond proceeds reported in funds', 'ch4'), ('Principal repayment reduces debt', 'ch5'),
       ('Increase in compensated absences', 'ch6'), ('Property-tax timing difference', 'ch7'), ('Internal Service Fund current-year loss', 'ch8'), ('Change in governmental activities net position', 'ch_tot')]
S.append({'id': 's4', 'title': 'Part 4: Reconcile the change in fund balances to the change in net position', 'intro': 'Governmental funds to governmental activities, operating statement. Enter reductions as negative numbers.', 'cols': ['Reconciliation item', 'Amount'],
          'rows': [{'label': l, 'cells': [{'k': k, 'a': float(find('Change Reconciliation', l)[1])}]} for l, k in chl]})
# S5 functional expense conversion
fe = [r for r in rows('Functional Expense Conversion', 5)]
fr = []
for i, r in enumerate(fe):
    cells = []
    for j in range(1, 7):
        v = float(r[j] or 0); cells.append({'k': 'fe%d_%d' % (i, j), 'a': v, 'graded': (v != 0 or j in (1, 6))})
    fr.append({'label': r[0], 'cells': cells})
S.append({'id': 's5', 'title': 'Part 5: Functional expense conversion', 'intro': 'Trace each function from fund expenditures to government-wide expenses. Remove capitalized equipment, then add depreciation, the increase in compensated absences and the allocated Motor Pool loss. Leave a cell blank if it is zero.',
          'cols': ['Function', 'Fund expenditures', 'Remove capital outlay', 'Add depreciation', 'Add compensated absences', 'Add ISF loss', 'Government-wide expense'], 'rows': fr})
# S6 SNP
sn = {r[0]: r for r in rows('Statement of Net Position', 5)}
S.append({'id': 's6', 'title': 'Part 6: Statement of Net Position (summary)', 'intro': 'June 30, 2027. Governmental activities include the Internal Service Fund; business-type activities are the Water Utility.', 'cols': ['Item', 'Governmental activities', 'Business-type activities', 'Total'],
          'rows': [{'label': l, 'cells': [{'k': 'snp%d_%d' % (i, j), 'a': float(sn[l][j])} for j in (1, 2, 3)]} for i, l in enumerate(['Net position', 'Beginning net position', 'Change in net position'])]})
# S7 SOA
soa = {r[0]: r for r in rows('Statement of Activities', 5) if r[0]}
gf = ['General government', 'Public safety', 'Public works', 'Parks and recreation', 'Interest on long-term debt']
sr = []
for i, l in enumerate(gf):
    r = soa[l]; sr.append({'label': l, 'cells': [{'k': 'so_e%d' % i, 'a': float(r[1])}, {'g': float(r[2])}, {'k': 'so_n%d' % i, 'a': float(r[3])}]})
t = soa['Total governmental activities']
sr.append({'label': 'Total governmental activities', 'cells': [{'k': 'so_te', 'a': float(t[1])}, {'k': 'so_tp', 'a': float(t[2])}, {'k': 'so_tn', 'a': float(t[3])}]})
for i, l in enumerate(['General revenues - property taxes', 'General revenues - sales taxes', 'General revenues - investment earnings']):
    sr.append({'label': l, 'cells': [None, None, {'k': 'so_g%d' % i, 'a': float(soa[l][1] if soa[l][1] is not None else soa[l][3])}]})
sr.append({'label': 'Change in governmental activities net position', 'cells': [None, None, {'k': 'so_cg', 'a': float(soa['Change in governmental activities net position'][3] if soa['Change in governmental activities net position'][3] is not None else soa['Change in governmental activities net position'][1])}]})
w = soa['Water Utility']
sr.append({'label': 'Water Utility (business-type activity)', 'cells': [{'k': 'so_we', 'a': float(w[1])}, {'g': float(w[2])}, {'k': 'so_wn', 'a': float(w[3])}]})
sr.append({'label': 'General revenue - investment earnings (business-type)', 'cells': [None, None, {'k': 'so_wi', 'a': float(soa['General revenue - investment earnings'][1] if soa['General revenue - investment earnings'][1] is not None else soa['General revenue - investment earnings'][3])}]})
cb = soa['Change in business-type activities net position']; sr.append({'label': 'Change in business-type activities net position', 'cells': [None, None, {'k': 'so_cb', 'a': float(cb[3] if cb[3] is not None else cb[1])}]})
ct = soa['Change in total primary government net position']; sr.append({'label': 'Change in total primary government net position', 'cells': [None, None, {'k': 'so_ct', 'a': float(ct[3] if ct[3] is not None else ct[1])}]})
S.append({'id': 's7', 'title': 'Part 7: Statement of Activities', 'intro': 'Fiscal year ended June 30, 2027. Program revenues by function are given. Expenses come from Part 5. Net (expense) revenue is program revenues minus expenses; enter net expenses as negative numbers.', 'cols': ['Function / program', 'Expenses', 'Program revenues (given)', 'Net (expense) revenue or amount'], 'rows': sr})
S.append({'id': 's8', 'title': 'Part 8: The custodial fund', 'intro': 'Answer the question below.', 'cols': [], 'rows': [], 'choice': {'k': 'cus', 'q': 'How is the Regional Tax Collection Custodial Fund reported in the government-wide statements?',
          'opts': ['It is included in governmental activities', 'It is included in business-type activities', 'It is excluded: fiduciary resources held for others are reported only in the fund statements', 'It is included only in the total column'], 'a': 2}})
TIES = [
 {'l': 'Fund balances by fund add to total governmental fund balances (Part 1)', 'a': ['fb_gf', 'fb_srf', 'fb_cpf', 'fb_dsf'], 'b': ['fb_tot']},
 {'l': 'Fund balance changes by fund add to the change in total governmental fund balances (Part 2)', 'a': ['cf_gf', 'cf_srf', 'cf_cpf', 'cf_dsf'], 'b': ['cf_tot']},
 {'l': 'Net position reconciliation lines add to governmental activities net position (Part 3)', 'a': ['np1', 'np2', 'np3', 'np4', 'np5', 'np6'], 'b': ['np_tot']},
 {'l': 'Reconciled net position equals the Statement of Net Position, governmental activities (Parts 3 and 6)', 'a': ['np_tot'], 'b': ['snp0_1']},
 {'l': 'Change reconciliation lines add to the change in governmental activities net position (Part 4)', 'a': ['ch1', 'ch2', 'ch3', 'ch4', 'ch5', 'ch6', 'ch7', 'ch8'], 'b': ['ch_tot']},
 {'l': 'Reconciled change equals the Statement of Activities change, governmental activities (Parts 4 and 7)', 'a': ['ch_tot'], 'b': ['so_cg']},
 {'l': 'Statement of Net Position change equals the Statement of Activities change, governmental activities (Parts 6 and 7)', 'a': ['snp2_1'], 'b': ['so_cg']},
 {'l': 'Government-wide expenses in Part 5 add to the Statement of Activities total expenses (Parts 5 and 7)', 'a': ['fe0_6', 'fe1_6', 'fe2_6', 'fe3_6', 'fe4_6'], 'b': ['so_te']}]
case = {'id': 'a5', 'title': 'Assignment 5: Government-Wide Conversion', 'facts': facts, 'sections': S, 'ties': TIES}
# verify every answer is numeric & ties hold on the key
ans = {}
for s in S:
    for r in s['rows']:
        for c in r['cells']:
            if c and 'k' in c: ans[c['k']] = c['a']
for t in TIES:
    a = sum(ans[k] for k in t['a']); b = sum(ans[k] for k in t['b'])
    print('TIE', 'OK ' if abs(a - b) < 1 else 'BAD', t['l'][:70], a, b)
print('cells', len(ans), 'graded', sum(1 for s in S for r in s['rows'] for c in r['cells'] if c and 'k' in c and c.get('graded', True)))
open('/home/claude/build/v2/a5.js', 'w').write('window.CASE5 = ' + json.dumps(case, separators=(',', ':')) + ';\n')
